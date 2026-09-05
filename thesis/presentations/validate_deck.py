"""Check a defence deck for the layout defects that a build script cannot see.

Every one of these checks exists because the fault it looks for was actually shipped in one of these
decks and found only by rendering the slide to an image and looking at it. Rendering is slow and
needs PowerPoint; these checks are fast and need neither, so they run first and catch the same faults
earlier.

  off-slide          a shape whose box leaves the 13.333 by 7.5 inch canvas. Caused by a column
                     positioned from the left when it should have been sized to the remaining width
  invisible text     a run whose colour matches the fill immediately behind it. Caused by using a
                     dark palette colour for a card title on a dark card
  collision          two runs of text whose ink overlaps. A card behind text is fine and is not
                     flagged; only the estimated ink of one run against another is compared, and the
                     ink is allowed to leave its declared box, because a top-anchored frame is not a
                     clip region and these decks place every element by an explicit y
  no notes           a slide with no speaker notes. Every slide in these decks is meant to be
                     defensible aloud

There was a fifth check, overset text, which reported any box whose text needed more height than the
box declared. It fired 32 times across the seven earliest decks and every one of them rendered
correctly: PowerPoint lays a top-anchored frame out from the top and simply continues past the
declared bottom, so the overflow landed in the blank space the layout had left for it. Sizing the
boxes up to satisfy the check turned 21 of those into collisions that were not there either. The
check was measuring a property these decks never maintained and no reader can see, so it is gone.
Overset is a collision or it is nothing.

The ink estimate is not a layout engine. Its glyph advance was measured off rendered slides rather
than assumed: 322 characters over 5 lines in a 5.19 inch box, and 358 over 11 lines in a 2.59 inch
box, both at 12.5 pt Calibri, give 0.446 and 0.452 em. The constant used to be 0.50, which is about
11 per cent too wide and was the reason two collisions were reported that do not happen.

Run:  brainsafe_env/Scripts/python.exe thesis/presentations/validate_deck.py
      brainsafe_env/Scripts/python.exe thesis/presentations/validate_deck.py chapter09_defence.pptx
Exit: 0 if every deck passes, 1 otherwise.
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation

from pptx.util import Emu

HERE = Path(__file__).resolve().parent
EMU_IN = 914400.0
W_IN, H_IN = 13.333, 7.5
TOL = 0.02                 # inches of slack at the slide edge
GLYPH = {"Cambria": 0.47, "Calibri": 0.45}   # measured off rendered slides, see the note above
# A line box is about 1.22 times the point size, but most of the extra is leading above and below
# the glyphs. Collisions are about ink, so the visible extent of a line is nearer one em: using the
# line box instead put a 40 pt display number 0.10 in taller than it draws and reported every
# value-and-label pair in the deck as a collision.
INK_FACTOR = 1.00



def _rgb(color):
    try:
        return tuple(color.rgb)
    except Exception:                                      # theme colour or inherited
        return None


def _fill_rgb(shape):
    try:
        f = shape.fill
        if f.type is not None and f.type == 1:             # MSO_FILL.SOLID
            return _rgb(f.fore_color)
    except Exception:
        pass
    return None


def _bg_rgb(slide):
    try:
        f = slide.background.fill
        if f.type == 1:
            return _rgb(f.fore_color)
    except Exception:
        pass
    return None


def _box(sh):
    return (sh.left / EMU_IN, sh.top / EMU_IN,
            (sh.left + sh.width) / EMU_IN, (sh.top + sh.height) / EMU_IN)


def _overlap(a, b):
    return not (a[2] <= b[0] or b[2] <= a[0] or a[3] <= b[1] or b[3] <= a[1])


def _near(c1, c2, tol=48):
    """Are two colours close enough that text in one is unreadable on the other?"""
    if c1 is None or c2 is None:
        return False
    return sum(abs(x - y) for x, y in zip(c1, c2)) < tol


def _ink_height(sh) -> float:
    """Inches of ink a text frame draws, summed over paragraphs.

    An estimate, not a layout engine: the measured average glyph advance for each face, and one em
    of vertical extent per line. Close enough to catch a block that runs into the one below it, and
    deliberately quiet about a few points either way.

    The space after the final paragraph is excluded because nothing is drawn in it.
    """
    tf = sh.text_frame
    width_in = sh.width / EMU_IN
    paras = list(tf.paragraphs)
    total = 0.0
    for k, p in enumerate(paras):
        runs = [r for r in p.runs if r.text]
        if not runs:
            total += 0.10
            continue
        size = max((r.font.size.pt if r.font.size else 12) for r in runs)
        face = runs[0].font.name or "Calibri"
        chars = sum(len(r.text) for r in runs)
        per_line = max(1, int(width_in * 72 / (size * GLYPH.get(face, 0.50))))
        # An explicit newline forces a break, but the text either side of it is short, so the
        # wrapped estimate and the break count are two views of the same lines rather than two
        # sets to be added. Take the larger.
        breaks = sum(r.text.count("\n") for r in runs)
        lines = max(1, -(-chars // per_line), breaks + 1)
        spacing = float(p.line_spacing) if isinstance(p.line_spacing, float) else 1.0
        total += lines * size * INK_FACTOR * spacing / 72.0
        if k < len(paras) - 1:
            total += (p.space_after.pt if p.space_after else 0) / 72.0
    return total


def validate(path: Path) -> list[str]:
    prs = Presentation(str(path))
    problems: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        bg = _bg_rgb(slide)
        cards = []                                          # (box, fill) of non-text shapes
        texts = []                                          # (box, shape) of shapes carrying text

        for sh in slide.shapes:
            if sh.left is None or sh.top is None:
                continue
            box = _box(sh)
            if (box[0] < -TOL or box[1] < -TOL
                    or box[2] > W_IN + TOL or box[3] > H_IN + TOL):
                problems.append(
                    f"slide {i}: off-slide shape at "
                    f"({box[0]:.2f}, {box[1]:.2f}) to ({box[2]:.2f}, {box[3]:.2f})")
            has_text = sh.has_text_frame and sh.text_frame.text.strip()
            fill = _fill_rgb(sh)
            if has_text:
                texts.append((box, sh))
            if fill is not None:
                cards.append((box, fill))

        for box, sh in texts:
            behind = bg
            for cbox, cfill in cards:
                if cbox != box and _overlap(box, cbox):
                    behind = cfill
            own = _fill_rgb(sh)
            if own is not None:
                behind = own
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if not r.text.strip():
                        continue
                    if _near(_rgb(r.font.color), behind):
                        problems.append(
                            f"slide {i}: text \"{r.text[:38]}\" is the colour of what is behind it")
                        break

        # Compare the ink, not the container, and let the ink leave the container.
        #
        # An earlier version of this file also reported any box whose text needed more height than
        # the box declared. It fired 32 times across the seven earliest decks and every one of them
        # rendered correctly, because a top-anchored text frame in PowerPoint is not a clip region:
        # the text is laid out from the top and simply continues past the declared bottom. These
        # decks position each element by an explicit y, so a box height is decoration and reporting
        # it as a fault was reporting on a property the decks never maintained and no reader can
        # see. Worse, sizing the boxes up to satisfy it turned 21 of those into collisions that were
        # not there either.
        #
        # What does matter is where the text ends up, so the estimate is no longer clipped to the
        # box: a top-anchored frame's ink runs from its top for as long as its text needs, and the
        # only complaint left is that the ink reaches another run. Overset is not a separate fault,
        # it is a collision or it is nothing.
        ink = []
        for box, sh in texts:
            need = _ink_height(sh)
            top = box[1]
            if sh.text_frame.vertical_anchor is not None and str(
                    sh.text_frame.vertical_anchor).startswith("MIDDLE"):
                need = min(need, box[3] - box[1])
                top = box[1] + max(0.0, (box[3] - box[1] - need) / 2)
            ink.append(((box[0], top, box[2], top + need), sh))

        for a in range(len(ink)):
            for b in range(a + 1, len(ink)):
                box_a, sh_a = ink[a]
                box_b, sh_b = ink[b]
                if not _overlap(box_a, box_b):
                    continue
                ov = (min(box_a[2], box_b[2]) - max(box_a[0], box_b[0])) * \
                     (min(box_a[3], box_b[3]) - max(box_a[1], box_b[1]))
                if ov > 0.06:
                    problems.append(
                        f"slide {i}: text overlaps by {ov:.2f} sq in: "
                        f"\"{sh_a.text_frame.text[:26].strip()}\" and "
                        f"\"{sh_b.text_frame.text[:26].strip()}\"")

        if not slide.has_notes_slide or not slide.notes_slide.notes_text_frame.text.strip():
            problems.append(f"slide {i}: no speaker notes")

    return problems


def main() -> int:
    names = sys.argv[1:] or sorted(p.name for p in HERE.glob("*.pptx"))
    bad = 0
    for name in names:
        path = HERE / name
        problems = validate(path)
        print(f"\n{name}: {len(Presentation(str(path)).slides._sldIdLst)} slides, "
              f"{len(problems)} problem(s)")
        for p in problems:
            print(f"  {p}")
        bad += len(problems)
    print(f"\n{'PASS' if not bad else 'FAIL'}: {bad} problem(s) across {len(names)} deck(s)")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
